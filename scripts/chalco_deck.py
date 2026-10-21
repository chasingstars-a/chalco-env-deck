# -*- coding: utf-8 -*-
"""chalco_deck —— 示例环保（EXAMPLE 环保节能）演示文稿样式库。

按「示例环保公司简介」模板的色调与 logo 复刻的一套 python-pptx 版式工具：
    * 16:9（13.333 x 7.5 in）
    * 品牌绿渐变（#2FCF9D → #33AC48 → #214B0D）+ 品牌蓝 #0056B0
    * 微软雅黑单一字族，标题加粗
    * 封面 / 目录 / 章节页 / 内容页 / 结尾页五种页面骨架 + 常用内容块

用法（在工作目录里写自己的 build 脚本）：
    import sys; sys.path.insert(0, r"<skill>\\scripts")
    from chalco_deck import Deck, C, FONT
    d = Deck()
    d.cover(title="公司简介", location="中国 · 北京", date="2026.10")
    d.toc(["公司简介", "科技立企", "业务能力", "商业模式"])
    d.section("01", "公司简介", en="COMPANY PROFILE", photo=asset("photo_leaf.jpg"))
    p = d.page("公司概况")
    p.bullets(["要点一", "要点二"], box=(0.95, 1.4, 6.3, 4.6))
    p.image(asset("photo_leaf.jpg"), box=(7.4, 1.4, 5.2, 4.6), radius=0.12)
    d.closing()
    d.save("output.pptx")

所有坐标单位都是英寸；颜色常量见 C。
"""
from __future__ import annotations

import hashlib
import os
import re
import tempfile
from pathlib import Path
from typing import Sequence

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

# --------------------------------------------------------------------------- #
# 基本常量
# --------------------------------------------------------------------------- #
SW, SH = 13.333, 7.5          # 幻灯片尺寸（英寸）
EMU = 914400                   # 1 英寸
FONT = "微软雅黑"

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"


def asset(name: str) -> str:
    """返回 skill assets 目录下素材的绝对路径。"""
    path = ASSETS_DIR / name
    if not path.exists():
        raise FileNotFoundError(f"缺少素材：{path}")
    return str(path)


class C:
    """模板取色。注释里的数值来自对原模板矢量填充 / logo 像素的直接采样。"""

    BLUE = "0056B0"          # 中铝徽标蓝（logo 主体）
    BLUE_DEEP = "003A78"     # 深蓝（遮罩、阴影）
    BLUE_LIGHT = "B2D5FD"    # 目录卡片、渐变起始浅蓝
    BLUE_WASH = "EDF1FF"     # 极浅蓝底
    GREEN_DEEP = "214B0D"    # 字标渐变末端，最深绿
    GREEN = "33AC48"         # 主绿（波浪、色块、表格头）
    GREEN_MID = "459753"     # 中绿（箭头渐变末端）
    GREEN_LIGHT = "50C563"   # 浅绿
    GREEN_SOFT = "67BB3E"    # 偏黄的绿（曲头渐变）
    MINT = "2FCF9D"          # 薄荷绿（渐变起点、强调线、数字）
    TEAL = "44C8A2"          # 青绿（小标签）
    MINT_PALE = "ACE6D5"     # 极浅薄荷（底纹）
    INK = "404040"           # 正文 / 标题主色
    INK_DEEP = "222222"      # 更重的标题
    GRAY = "7F7F7F"          # 页码、次级文字
    GRAY_LIGHT = "BFBFBF"    # 章节数字
    LINE = "DCE6F0"          # 分隔线
    BG = "F1F4FA"            # 浅蓝白底
    WHITE = "FFFFFF"
    BAND = "3FA45C"          # 结尾页绿色横条


# 允许出现在成品里的颜色（自检脚本用）
PALETTE = {
    C.BLUE, C.BLUE_DEEP, C.BLUE_LIGHT, C.BLUE_WASH, C.GREEN_DEEP, C.GREEN,
    C.GREEN_MID, C.GREEN_LIGHT, C.GREEN_SOFT, C.MINT, C.TEAL, C.MINT_PALE,
    C.INK, C.INK_DEEP, C.GRAY, C.GRAY_LIGHT, C.LINE, C.BG, C.WHITE, C.BAND,
    # 渐变中间色 / 辅助色
    "15329B", "09E4D8", "1F4B8F", "2E8B4F", "F5FAF7", "E8F5EC", "F2F7FB",
    "A8CBF0", "D8E8F8", "9FC4EE", "DCEAF8", "C9DCF0", "C9D8E8", "F7FAFD", "606060",
    "4CB268", "2F8E4E", "000000",
}

CONTENT_BOX = (0.95, 1.30, 11.43, 5.55)   # 内容页默认可用区域 (x, y, w, h)


# --------------------------------------------------------------------------- #
# 颜色 / 效果小工具
# --------------------------------------------------------------------------- #
def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.upper())


def _spPr(shape):
    return shape._element.spPr


_FILL_TAGS = ("a:noFill", "a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill", "a:grpFill")


def _drop_fill(spPr):
    for tag in _FILL_TAGS:
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)


def _insert_fill(spPr, el):
    """把填充元素插到 a:ln 之前，保证 spPr 子元素顺序合法。"""
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(el)
    else:
        spPr.append(el)


def _insert_ln(spPr, ln):
    for tag in ("a:effectLst", "a:effectDag", "a:scene3d", "a:sp3d"):
        el = spPr.find(qn(tag))
        if el is not None:
            el.addprevious(ln)
            return
    spPr.append(ln)


def fill_solid(shape, color: str, alpha: float | None = None):
    spPr = _spPr(shape)
    _drop_fill(spPr)
    inner = f'<a:srgbClr val="{color}"/>'
    if alpha is not None and alpha < 1:
        inner = f'<a:srgbClr val="{color}"><a:alpha val="{int(alpha * 100000)}"/></a:srgbClr>'
    _insert_fill(spPr, parse_xml(f'<a:solidFill {nsdecls("a")}>{inner}</a:solidFill>'))
    return shape


def fill_grad(shape, stops: Sequence[tuple], angle: float = 90.0):
    """线性渐变。stops = [(pos0~1, 'RRGGBB', alpha 或 None), ...]。"""
    spPr = _spPr(shape)
    _drop_fill(spPr)
    parts = []
    for pos, color, alpha in stops:
        if alpha is not None and alpha < 1:
            inner = f'<a:srgbClr val="{color}"><a:alpha val="{int(alpha * 100000)}"/></a:srgbClr>'
        else:
            inner = f'<a:srgbClr val="{color}"/>'
        parts.append(f'<a:gs pos="{int(round(pos * 100000))}">{inner}</a:gs>')
    xml = (
        f'<a:gradFill {nsdecls("a")} rotWithShape="1"><a:gsLst>{"".join(parts)}</a:gsLst>'
        f'<a:lin ang="{int(round(angle * 60000))}" scaled="0"/></a:gradFill>'
    )
    _insert_fill(spPr, parse_xml(xml))
    return shape


def no_line(shape):
    spPr = _spPr(shape)
    for el in spPr.findall(qn("a:ln")):
        spPr.remove(el)
    _insert_ln(spPr, parse_xml(f'<a:ln {nsdecls("a")} w="0"><a:noFill/></a:ln>'))
    return shape


def set_line(shape, color: str, width_pt: float = 1.0, alpha: float | None = None):
    spPr = _spPr(shape)
    for el in spPr.findall(qn("a:ln")):
        spPr.remove(el)
    if alpha is not None and alpha < 1:
        inner = f'<a:srgbClr val="{color}"><a:alpha val="{int(alpha * 100000)}"/></a:srgbClr>'
    else:
        inner = f'<a:srgbClr val="{color}"/>'
    ln = parse_xml(
        f'<a:ln {nsdecls("a")} w="{int(width_pt * 12700)}" cap="flat" cmpd="sng" algn="ctr">'
        f"<a:solidFill>{inner}</a:solidFill><a:prstDash val=\"solid\"/></a:ln>"
    )
    _insert_ln(spPr, ln)
    return shape


def shadow(shape, blur_pt: float = 16, color: str = "15329B", alpha: float = 0.18,
           dist_pt: float = 4, dir_deg: float = 90, glow: bool = False):
    """外阴影；glow=True 时改成模板里那种四向柔光。"""
    spPr = _spPr(shape)
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    if glow:
        eff = (
            f'<a:outerShdw blurRad="{int(blur_pt * 12700)}" sx="102000" sy="102000" '
            f'algn="ctr" rotWithShape="0"><a:srgbClr val="{color}">'
            f'<a:alpha val="{int(alpha * 100000)}"/></a:srgbClr></a:outerShdw>'
        )
    else:
        eff = (
            f'<a:outerShdw blurRad="{int(blur_pt * 12700)}" dist="{int(dist_pt * 12700)}" '
            f'dir="{int(dir_deg * 60000)}" rotWithShape="0"><a:srgbClr val="{color}">'
            f'<a:alpha val="{int(alpha * 100000)}"/></a:srgbClr></a:outerShdw>'
        )
    spPr.append(parse_xml(f'<a:effectLst {nsdecls("a")}>{eff}</a:effectLst>'))
    return shape


# --------------------------------------------------------------------------- #
# 文字
# --------------------------------------------------------------------------- #
def _style_run(run, size, color, bold=False, font=FONT, spacing=None, italic=False,
               text_shadow=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    if rPr.find(qn("a:ea")) is None:
        rPr.append(parse_xml(f'<a:ea {nsdecls("a")} typeface="{font}"/>'))
    if spacing:
        rPr.set("spc", str(int(round(spacing * 100))))
    if text_shadow:
        el = parse_xml(
            f'<a:effectLst {nsdecls("a")}><a:outerShdw blurRad="63500" dist="25400" '
            f'dir="5400000" rotWithShape="0"><a:srgbClr val="000000">'
            f'<a:alpha val="35000"/></a:srgbClr></a:outerShdw></a:effectLst>'
        )
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addprevious(el)
        else:
            rPr.append(el)
    return run


_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
_ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def text(slide, box, content, size=16, color=C.INK, bold=False, align="l", anchor="t",
         spacing=None, line_spacing=1.25, space_after=0, font=FONT, name=None,
         rotation=None, text_shadow=False, wrap=True):
    """在 box=(x,y,w,h) 里放文字。

    content 可以是：
        "一行字"            —— 含 "\\n" 时按段落拆开
        ["第一段", "第二段"]
        [[("普通", {}), ("重点", {"color": C.GREEN, "bold": True})], "第二段"]
    """
    x, y, w, h = box
    tb = slide.shapes.add_textbox(Emu(int(x * EMU)), Emu(int(y * EMU)),
                                  Emu(int(w * EMU)), Emu(int(h * EMU)))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR[anchor]
    if isinstance(content, str):
        # 用独立段落承载换行：渲染器不会把 <a:br/> 之后的行按段首对齐
        paragraphs = content.split("\n")
    elif isinstance(content, (list, tuple)):
        paragraphs = content
    else:
        paragraphs = [str(content)]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, (list, tuple)) else [(para, {})]
        for raw in runs:
            if isinstance(raw, str):
                chunk, opts = raw, {}
            else:
                chunk, opts = raw[0], dict(raw[1] if len(raw) > 1 else {})
            r = p.add_run()
            r.text = chunk
            _style_run(r, opts.get("size", size), opts.get("color", color),
                       opts.get("bold", bold), opts.get("font", font),
                       opts.get("spacing", spacing), opts.get("italic", False),
                       opts.get("text_shadow", text_shadow))
    if rotation:
        tb.rotation = rotation
    return tb


def bullets(slide, box, items, size=16, color=C.INK, bold=False, marker="■",
            marker_color=C.GREEN, line_spacing=1.35, space_after=8, marL=0.28,
            align="l", anchor="t", marker_size=0.55, name=None):
    """绿色方块项目符号列表（原生 buChar，换行自动对齐）。"""
    x, y, w, h = box
    tb = slide.shapes.add_textbox(Emu(int(x * EMU)), Emu(int(y * EMU)),
                                  Emu(int(w * EMU)), Emu(int(h * EMU)))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR[anchor]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        p.line_spacing = line_spacing
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(marL * EMU)))
        pPr.set("indent", str(int(-marL * EMU)))
        pPr.append(parse_xml(
            f'<a:buClr {nsdecls("a")}><a:srgbClr val="{marker_color}"/></a:buClr>'))
        pPr.append(parse_xml(
            f'<a:buSzPct {nsdecls("a")} val="{int(marker_size * 100000)}"/>'))
        pPr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="Arial"/>'))
        pPr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="{marker}"/>'))
        runs = item if isinstance(item, (list, tuple)) else [(item, {})]
        for raw in runs:
            if isinstance(raw, str):
                chunk, opts = raw, {}
            else:
                chunk, opts = raw[0], dict(raw[1] if len(raw) > 1 else {})
            r = p.add_run()
            r.text = chunk
            _style_run(r, opts.get("size", size), opts.get("color", color),
                       opts.get("bold", bold), FONT, opts.get("spacing"))
    return tb


# --------------------------------------------------------------------------- #
# 文字度量（自动收缩字号 / 溢出检测共用）
# --------------------------------------------------------------------------- #
_FONT_FILES = {
    False: (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhl.ttc"),
    True: (r"C:\Windows\Fonts\msyhbd.ttc", r"C:\Windows\Fonts\msyhbd.ttc"),
}
_font_cache: dict = {}


def _pil_font(size_pt: float, bold: bool = False):
    key = (round(size_pt * 4), bold)
    if key not in _font_cache:
        candidates = _FONT_FILES[bool(bold)]
        path = next((p for p in candidates if os.path.exists(p)), None)
        _font_cache[key] = (ImageFont.truetype(path, max(6, int(round(size_pt * 96 / 72))))
                            if path else ImageFont.load_default())
    return _font_cache[key]


def _is_wide(ch: str) -> bool:
    o = ord(ch)
    return (0x2E80 <= o <= 0x9FFF or 0xF900 <= o <= 0xFAFF
            or 0xFF00 <= o <= 0xFFEF or 0x3000 <= o <= 0x303F)


def wrap_lines(text_str: str, width_in: float, size_pt: float, bold: bool = False) -> list[str]:
    """按像素宽度粗略折行（中文逐字、英文尽量整词）。"""
    font = _pil_font(size_pt, bold)
    limit = max(1.0, width_in * 96)
    out: list[str] = []
    for para in str(text_str).split("\n"):
        if not para:
            out.append("")
            continue
        cur, cur_w, token = "", 0.0, ""

        def flush():
            nonlocal cur, cur_w, token
            if not token:
                return
            w = font.getlength(token)
            if cur_w + w > limit and cur:
                out.append(cur)
                cur, cur_w = "", 0.0
            cur += token
            cur_w += w
            token = ""

        for ch in para:
            if _is_wide(ch) or ch == " ":
                flush()
                w = font.getlength(ch)
                if cur_w + w > limit and cur:
                    out.append(cur)
                    cur, cur_w = "", 0.0
                cur += ch
                cur_w += w
            else:
                token += ch
                if font.getlength(token) > limit and len(token) > 1:
                    token = token[:-1]
                    flush()
                    token = ch
        flush()
        out.append(cur)
    return out


def text_height(text_str: str, width_in: float, size_pt: float, bold: bool = False,
                line_spacing: float = 1.25) -> float:
    n = len(wrap_lines(text_str, width_in, size_pt, bold))
    return n * size_pt * 1.2 * line_spacing / 72.0


def fit_size(text_str: str, box, size: float, bold: bool = False, line_spacing: float = 1.25,
             min_size: float = 8.0, step: float = 1.0) -> float:
    """在不超框的前提下挑一个能放下的字号。box 可以是 (w, h) 或 (x, y, w, h)。"""
    w, h = (box[0], box[1]) if len(box) == 2 else (box[2], box[3])
    cur = size
    while cur > min_size and text_height(text_str, w, cur, bold, line_spacing) > h:
        cur -= step
    return max(cur, min_size)


# --------------------------------------------------------------------------- #
# 图形 / 图片
# --------------------------------------------------------------------------- #
def rect(slide, box, color=None, radius=None, grad=None, angle=90.0, line=None,
         line_w=1.0, alpha=None, shadow_kw=None, name=None, shape=None):
    """矩形 / 圆角矩形 / 任意预设形状。box=(x,y,w,h)。"""
    x, y, w, h = box
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, Emu(int(x * EMU)), Emu(int(y * EMU)),
                                 Emu(int(w * EMU)), Emu(int(h * EMU)))
    if name:
        shp.name = name
    # 去掉主题样式引用：否则未显式设阴影时，主题会给形状叠一层立体阴影
    for style in shp._element.findall(qn("p:style")):
        shp._element.remove(style)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE and radius:
        shp.adjustments[0] = min(0.5, float(radius) / max(0.01, min(w, h)))
    if grad:
        fill_grad(shp, grad, angle)
    elif color:
        fill_solid(shp, color, alpha)
    else:
        spPr = _spPr(shp)
        _drop_fill(spPr)
        spPr.append(parse_xml(f'<a:noFill {nsdecls("a")}/>'))
    if line:
        set_line(shp, line, line_w)
    else:
        no_line(shp)
    if shadow_kw:
        shadow(shp, **shadow_kw)
    else:
        _spPr(shp).append(parse_xml(f'<a:effectLst {nsdecls("a")}/>'))
    shp.text_frame.word_wrap = True
    return shp


def _chev(slide, box, color, grad, angle, adj, alpha, name):
    shp = rect(slide, box, color=color, grad=grad, angle=angle, alpha=alpha,
               shape=MSO_SHAPE.CHEVRON, name=name)
    try:
        shp.adjustments[0] = adj
    except Exception:
        pass
    return shp


def chevron(slide, box, color=None, grad=None, angle=0.0, adj=0.5, alpha=None, name=None):
    """V 形箭头；adj 控制箭头深度（0.2 尖、0.5 深）。"""
    return _chev(slide, box, color, grad, angle, adj, alpha, name)


def band(slide, box, grad, tip=0.32, name="横幅"):
    """双尖头色带（左右各一个箭头尖），自由多边形绘制，缩放不变形。"""
    x, y, w, h = box
    pts = [(x + tip, y), (x + w - tip, y), (x + w, y + h / 2),
           (x + w - tip, y + h), (x + tip, y + h), (x, y + h / 2)]
    ff = slide.shapes.build_freeform(Emu(int(pts[0][0] * EMU)), Emu(int(pts[0][1] * EMU)))
    ff.add_line_segments([(Emu(int(px * EMU)), Emu(int(py * EMU))) for px, py in pts[1:]],
                         close=True)
    shp = ff.convert_to_shape()
    fill_grad(shp, grad, 0)
    no_line(shp)
    shp.name = name
    return shp


def _corner_mask(size, radius_px, corners=("tl", "tr", "bl", "br")):
    mask = Image.new("L", size, 0)
    d = ImageDraw.Draw(mask)
    r = int(max(0, min(radius_px, min(size) / 2)))
    d.rounded_rectangle([0, 0, size[0] - 1, size[1] - 1], radius=r, fill=255)
    if r and len(corners) < 4:
        if "tl" not in corners:
            d.rectangle([0, 0, r, r], fill=255)
        if "tr" not in corners:
            d.rectangle([size[0] - r - 1, 0, size[0] - 1, r], fill=255)
        if "bl" not in corners:
            d.rectangle([0, size[1] - r - 1, r, size[1] - 1], fill=255)
        if "br" not in corners:
            d.rectangle([size[0] - r - 1, size[1] - r - 1, size[0] - 1, size[1] - 1], fill=255)
    return mask


_CACHE_DIR = Path(tempfile.gettempdir()) / "chalco_deck_cache"


def _prepare_image(path: str, w_in: float, h_in: float, radius_in: float = 0.0,
                   focus=(0.5, 0.5), dpi: int = 160, corners=("tl", "tr", "bl", "br")) -> str:
    """按目标框裁切/缩放（object-fit: cover），可选圆角，结果缓存复用。

    带透明通道的素材（徽标、字标）保留为 PNG，不会被压成黑底。
    """
    _CACHE_DIR.mkdir(parents=True, exist_ok=True)
    w_px = max(2, int(round(w_in * dpi)))
    h_px = max(2, int(round(h_in * dpi)))
    key = f"{os.path.abspath(path)}|{w_px}x{h_px}|{radius_in}|{focus}|{corners}".encode()
    tag = hashlib.sha1(key).hexdigest()[:16]
    for candidate in _CACHE_DIR.glob(f"{tag}.*"):
        return str(candidate)
    im = Image.open(path).convert("RGBA")
    src_ar, dst_ar = im.width / im.height, w_px / h_px
    if src_ar > dst_ar:
        new_w = max(1, int(round(im.height * dst_ar)))
        left = int(round((im.width - new_w) * focus[0]))
        im = im.crop((left, 0, left + new_w, im.height))
    else:
        new_h = max(1, int(round(im.width / dst_ar)))
        top = int(round((im.height - new_h) * focus[1]))
        im = im.crop((0, top, im.width, top + new_h))
    im = im.resize((w_px, h_px), Image.LANCZOS)
    has_alpha = im.getextrema()[3][0] < 250
    if radius_in > 0:
        im.putalpha(_corner_mask((w_px, h_px), radius_in * dpi, corners))
        has_alpha = True
    if has_alpha:
        out = _CACHE_DIR / f"{tag}.png"
        im.save(out, "PNG", optimize=True)
    else:
        out = _CACHE_DIR / f"{tag}.jpg"
        im.convert("RGB").save(out, "JPEG", quality=90, optimize=True, progressive=True)
    return str(out)


def pic(slide, box, path, radius=0.0, focus=(0.5, 0.5), dpi=160, shadow_kw=None,
        corners=("tl", "tr", "bl", "br"), name=None):
    """放一张图：自动裁切填满 box，radius>0 时四角变圆。"""
    x, y, w, h = box
    prepared = _prepare_image(str(path), w, h, radius, focus, dpi, corners)
    p = slide.shapes.add_picture(prepared, Emu(int(x * EMU)), Emu(int(y * EMU)),
                                 Emu(int(w * EMU)), Emu(int(h * EMU)))
    p.name = name or f"图片 {os.path.basename(path)}"
    if shadow_kw:
        shadow(p, **shadow_kw)
    return p


def bg(slide, path=None, color=C.BG):
    """整页背景：图片铺满或纯色。"""
    if path:
        return pic(slide, (0, 0, SW, SH), path, dpi=200, name="背景")
    return rect(slide, (0, 0, SW, SH), color=color, name="背景")


def page_number(slide, n, color=C.GRAY, size=14):
    return text(slide, (12.70, 7.075, 0.45, 0.26), str(n), size=size, color=color,
                align="r", anchor="m", name="页码")


def logo(slide, x=11.91, y=0.374, w=0.687, white=False):
    name = "logo_emblem_white.png" if white else "logo_emblem.png"
    h = w * 209 / 205
    return pic(slide, (x, y, w, h), asset(name), name="logo")


def tiny_chevrons(slide, x, y, n=4, h=0.238, w=0.185, gap=0.0235,
                  colors=(C.MINT_PALE, C.GREEN), alpha=0.95):
    """标题左侧那排小箭头（第 1 个是五边形，后面 V 形）。"""
    for i in range(n):
        cx = x + i * (w + gap)
        stops = [(0, colors[0], alpha), (1, colors[1], alpha)]
        if i == 0:
            rect(slide, (cx, y, w * 0.84, h), grad=stops, angle=0,
                 shape=MSO_SHAPE.PENTAGON, name="小箭头")
        else:
            _chev(slide, (cx, y, w, h), None, stops, 0, 0.504, None, "小箭头")
    return None


def _cell_border(cell, side: str, color: str | None = None, width_pt: float = 0.75):
    """给表格单元格设边框。side ∈ L/R/T/B；color=None 表示无线。"""
    tcPr = cell._tc.get_or_add_tcPr()
    tag = f"a:ln{side}"
    for el in tcPr.findall(qn(tag)):
        tcPr.remove(el)
    if color is None:
        el = parse_xml(f'<{tag} {nsdecls("a")} w="0" cap="flat" cmpd="sng" algn="ctr">'
                       f"<a:noFill/></{tag}>")
    else:
        el = parse_xml(
            f'<{tag} {nsdecls("a")} w="{int(width_pt * 12700)}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></{tag}>')
    tcPr.insert(0, el)


# --------------------------------------------------------------------------- #
# 内容页
# --------------------------------------------------------------------------- #
class Page:
    """内容页：已带背景、标题、logo、页码；方法负责摆内容块。"""

    def __init__(self, slide, title, box=CONTENT_BOX):
        self.slide = slide
        self.title = title
        self.box = box

    # --- 基础 ---
    def text(self, content, box=None, **kw):
        return text(self.slide, box or self.box, content, **kw)

    def bullets(self, items, box=None, **kw):
        return bullets(self.slide, box or self.box, items, **kw)

    def rect(self, box, **kw):
        return rect(self.slide, box, **kw)

    def rule(self, box, color=C.MINT, grad=True):
        if grad:
            return rect(self.slide, box, grad=[(0, C.MINT, None), (1, C.GREEN, None)], angle=0)
        return rect(self.slide, box, color=color)

    def image(self, path, box=None, radius=0.10, caption=None, caption_size=11,
              focus=(0.5, 0.5), shadow_kw=None):
        box = box or self.box
        p = pic(self.slide, box, path, radius=radius, focus=focus, shadow_kw=shadow_kw)
        if caption:
            text(self.slide, (box[0], box[1] + box[3] + 0.06, box[2], 0.26),
                 caption, size=caption_size, color=C.GRAY, align="c")
        return p

    # --- 常用内容块 ---
    def card(self, box, title, lines=(), image=None, image_h=None, title_size=18,
             body_size=15, accent=True, center=True, body_color=C.INK,
             card_color=C.WHITE, shadow_kw=None):
        """白底圆角卡片：可选顶部图片 + 标题 + 渐变短线 + 正文；文字区整体垂直居中。"""
        x, y, w, h = box
        radius = 0.10
        rect(self.slide, box, color=card_color, radius=radius,
             shadow_kw=shadow_kw or dict(blur_pt=20, color="1F4B8F", alpha=0.20, glow=True))
        top = y
        if image:
            ih = image_h or min(h * 0.42, 2.3)
            pic(self.slide, (x, y, w, ih), image, radius=radius, dpi=170,
                corners=("tl", "tr"))
            rect(self.slide, (x, y + ih - 0.02, w, 0.04), color=card_color)
            top = y + ih
        pad = 0.22
        inner_w = w - 2 * pad
        lines = [str(v) for v in lines]
        body = "\n".join(lines)
        room = max(0.4, y + h - top - 0.34)
        reserved = (0.42 + (0.20 if accent else 0.0) + 0.10) if title else 0.0
        bsize = body_size
        while bsize > 10 and lines and \
                text_height(body, inner_w, bsize, line_spacing=1.45) > room - reserved:
            bsize -= 0.5
        body_h = text_height(body, inner_w, bsize, line_spacing=1.45) if lines else 0.0
        block_h = reserved + body_h
        if center:
            cursor = top + max(0.20, (y + h - top - block_h) / 2)
        else:
            cursor = top + 0.22
        if title:
            text(self.slide, (x + pad, cursor, inner_w, 0.42), title,
                 size=title_size, color=C.INK, bold=True,
                 align="c" if center else "l", anchor="m")
            cursor += 0.46
            if accent:
                bar_w = min(1.55, w * 0.52)
                bx = x + (w - bar_w) / 2 if center else x + pad
                rect(self.slide, (bx, cursor - 0.02, bar_w, 0.07),
                     grad=[(0, C.MINT, None), (1, C.GREEN, None)], angle=0, name="装饰线")
                cursor += 0.20
        if lines:
            text(self.slide, (x + pad, cursor, inner_w, max(0.3, body_h + 0.12)), lines,
                 size=bsize, color=body_color, align="c" if center else "l",
                 line_spacing=1.45, anchor="t")
        return box

    def cards(self, items, box=None, cols=None, gap=0.30, image_h=2.25, title_size=18,
              body_size=15):
        """一行多张卡片。items = [dict(title=, lines=[], image=), ...]。"""
        x, y, w, h = box or self.box
        n = len(items)
        cols = cols or n
        cw = (w - gap * (cols - 1)) / cols
        for i, item in enumerate(items):
            cx = x + i * (cw + gap)
            self.card((cx, y, cw, h), item.get("title", ""), item.get("lines", ()),
                      item.get("image"), image_h=image_h,
                      title_size=item.get("title_size", title_size),
                      body_size=item.get("body_size", body_size),
                      center=item.get("center", True))
        return self

    def kpis(self, items, box=None, gap=0.28, card=True, value_size=34, label_size=13):
        """指标卡。items = [(数值, 单位, 说明), ...]。"""
        x, y, w, h = box or self.box
        n = len(items)
        cw = (w - gap * (n - 1)) / n
        for i, item in enumerate(items):
            value, unit, note = (list(item) + ["", ""])[:3]
            bx = x + i * (cw + gap)
            if card:
                rect(self.slide, (bx, y, cw, h), color=C.WHITE, radius=0.10,
                     shadow_kw=dict(blur_pt=18, color="1F4B8F", alpha=0.18, glow=True))
            rect(self.slide, (bx + 0.28, y + 0.26, 0.42, 0.055),
                 grad=[(0, C.MINT, None), (1, C.GREEN, None)], angle=0)
            text(self.slide, (bx + 0.28, y + 0.44, cw - 0.56, 0.75),
                 [[(str(value), {"size": value_size, "bold": True, "color": C.GREEN})]
                  + ([(f" {unit}", {"size": 15, "bold": True, "color": C.GREEN})] if unit else [])],
                 anchor="m")
            text(self.slide, (bx + 0.28, y + 1.22, cw - 0.56, h - 1.42), note,
                 size=label_size, color=C.INK, line_spacing=1.3)
        return self

    def table(self, headers, rows, box=None, col_widths=None, size=12, header_size=12,
              first_col_bold=True, row_h=0.44, header_fill=C.GREEN, col_align=None):
        """品牌化表格：绿色表头 + 浅绿斑马纹 + 极浅行间线。

        col_align：每列对齐方式，如 ["c", "l", "l", "c"]；省略时首列左对齐、其余居中。
        """
        x, y, w, h = box or self.box
        n_rows, n_cols = len(rows) + 1, len(headers)
        row_h = min(row_h, h / n_rows)
        shape = self.slide.shapes.add_table(n_rows, n_cols, Emu(int(x * EMU)),
                                            Emu(int(y * EMU)), Emu(int(w * EMU)),
                                            Emu(int(row_h * n_rows)))
        tbl = shape.table
        tbl.first_row = True
        tbl.horz_banding = False
        tblPr = tbl._tbl.tblPr
        for el in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(el)
        tblPr.append(parse_xml(
            f'<a:tableStyleId {nsdecls("a")}>{{2D5ABB26-0587-4C30-8999-92F81FD0307C}}</a:tableStyleId>'))
        if col_widths:
            total = sum(col_widths)
            for j, cw in enumerate(col_widths):
                tbl.columns[j].width = Emu(int(w * cw / total * EMU))
        for r in range(n_rows):
            tbl.rows[r].height = Emu(int(row_h * EMU))
            for c in range(n_cols):
                cell = tbl.cell(r, c)
                cell.margin_left = Emu(91440)
                cell.margin_right = Emu(91440)
                cell.margin_top = Emu(30000)
                cell.margin_bottom = Emu(30000)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                if r == 0:
                    fill_hex = header_fill
                else:
                    fill_hex = "FFFFFF" if r % 2 else "F5FAF7"
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fill_hex)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                if col_align:
                    p.alignment = _ALIGN[col_align[c]]
                else:
                    p.alignment = PP_ALIGN.CENTER if (c > 0 or r == 0) else PP_ALIGN.LEFT
                p.line_spacing = 1.15
                for old in list(p.runs):
                    old._r.getparent().remove(old._r)
                value = headers[c] if r == 0 else rows[r - 1][c]
                run = p.add_run()
                run.text = str(value)
                if r == 0:
                    _style_run(run, header_size, C.WHITE, True)
                else:
                    _style_run(run, size, C.INK, bool(first_col_bold and c == 0))
                # 竖线用单元格自身底色，避免出现深浅不一的表格框线
                _cell_border(cell, "B", C.LINE, 0.75)
                _cell_border(cell, "T", fill_hex, 0.75)
                _cell_border(cell, "R", fill_hex, 0.75)
                _cell_border(cell, "L", fill_hex, 0.75)
        return shape

    def process(self, steps, box=None, gap=0.0, colors=None, text_size=12, title_size=14,
                pad_left=0.30, pad_right=0.18, adj=0.32):
        """横向箭头流程条。steps = [(标题, 说明), ...]。

        列数 ≥5 时把 pad_left/pad_right 调小（如 0.18/0.06），否则文字会被挤成两行。
        """
        x, y, w, h = box or self.box
        n = len(steps)
        cw = (w + gap * (n - 1)) / n
        shades = colors or [
            (C.MINT, C.GREEN_LIGHT), (C.GREEN_LIGHT, C.GREEN),
            (C.GREEN, C.GREEN_DEEP), (C.TEAL, C.GREEN_MID),
        ]
        for i, step in enumerate(steps):
            title, desc = (list(step) + [""])[:2]
            bx = x + i * (cw - gap)
            s, e = shades[i % len(shades)]
            shp = _chev(self.slide, (bx, y, cw, h), None,
                        [(0, s, None), (1, e, None)], 0, adj, None, "流程箭头")
            tf = shp.text_frame
            tf.word_wrap = True
            tf.margin_left = Emu(int(pad_left * EMU))
            tf.margin_right = Emu(int(pad_right * EMU))
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = title
            _style_run(run, title_size, C.WHITE, True)
            if desc:
                for line in str(desc).split("\n"):
                    p2 = tf.add_paragraph()
                    p2.alignment = PP_ALIGN.CENTER
                    run2 = p2.add_run()
                    run2.text = line
                    _style_run(run2, text_size, C.WHITE, False)
                    p2.line_spacing = 1.15
        return self

    def banner(self, title, box=None, grad_colors=((0, C.GREEN, None), (1, C.GREEN_MID, None)),
               size=20, tip=0.32):
        """绿色双尖头横幅（模板「三大板块」那条）。"""
        x, y, w, h = box or (0.95, 5.55, 11.43, 0.52)
        shp = band(self.slide, (x, y, w, h), list(grad_colors), tip=tip, name="横幅")
        text(self.slide, (x + tip + 0.2, y, w - 2 * tip - 0.4, h), title, size=size,
             color=C.WHITE, bold=True, align="c", anchor="m", spacing=1)
        return shp

    def tags(self, items, box=None, size=13, color=C.GREEN, text_color=C.WHITE,
             h=0.36, gap=0.18, filled=True):
        """小圆角标签，从左到右排列。"""
        x, y = (box or self.box)[:2]
        cx = x
        for item in items:
            w = max(0.9, 0.34 + len(str(item)) * size / 72 * 1.05)
            if filled:
                rect(self.slide, (cx, y, w, h), color=color, radius=h / 2)
                text(self.slide, (cx, y, w, h), str(item), size=size, color=text_color,
                     bold=True, align="c", anchor="m")
            else:
                rect(self.slide, (cx, y, w, h), color=C.WHITE, radius=h / 2,
                     line=color, line_w=1.0)
                text(self.slide, (cx, y, w, h), str(item), size=size, color=color,
                     bold=True, align="c", anchor="m")
            cx += w + gap
        return self

    def note(self, content, box=None, size=13, color=C.INK, fill="E8F5EC",
             bar_color=C.GREEN):
        """浅绿提示条。"""
        x, y, w, h = box or (0.95, 6.10, 11.43, 0.5)
        rect(self.slide, (x, y, w, h), color=fill, radius=0.08)
        rect(self.slide, (x, y, 0.07, h), color=bar_color)
        text(self.slide, (x + 0.24, y, w - 0.4, h), content, size=size, color=color,
             anchor="m", line_spacing=1.25)
        return self

    def timeline(self, items, box=None, size=13, title_size=14):
        """竖向时间线。items = [(时间, 标题, 说明), ...]。"""
        x, y, w, h = box or self.box
        n = len(items)
        step = h / n
        line_x = x + 1.16
        rect(self.slide, (line_x, y + 0.12, 0.035, h - 0.24), color=C.MINT_PALE)
        for i, item in enumerate(items):
            when, title, desc = (list(item) + ["", ""])[:3]
            cy = y + i * step
            rect(self.slide, (x, cy + 0.02, 1.02, 0.32), color=C.GREEN, radius=0.06)
            text(self.slide, (x, cy + 0.02, 1.02, 0.32), when, size=11, color=C.WHITE,
                 bold=True, align="c", anchor="m")
            rect(self.slide, (line_x - 0.035, cy + 0.13, 0.105, 0.105), color=C.WHITE,
                 shape=MSO_SHAPE.OVAL, line=C.GREEN, line_w=2)
            text(self.slide, (x + 1.40, cy, w - 1.40, 0.30), title, size=title_size,
                 color=C.INK, bold=True, anchor="m")
            if desc:
                text(self.slide, (x + 1.40, cy + 0.30, w - 1.40, step - 0.34), desc,
                     size=size, color=C.INK, line_spacing=1.25)
        return self

    def chart(self, categories, series, box=None, kind="bar", number_format=None,
              legend=True, size=11):
        """原生可编辑图表（柱/条形/折线）。series = [(名称, [数值...]), ...]。"""
        x, y, w, h = box or self.box
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        chart_type = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                      "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
                      "line": XL_CHART_TYPE.LINE_MARKERS}[kind]
        gframe = self.slide.shapes.add_chart(chart_type, Emu(int(x * EMU)), Emu(int(y * EMU)),
                                             Emu(int(w * EMU)), Emu(int(h * EMU)), data)
        chart = gframe.chart
        chart.has_legend = legend
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(size)
            chart.legend.font.name = FONT
        chart.font.size = Pt(size)
        chart.font.name = FONT
        colors = [C.GREEN, C.MINT, C.BLUE, C.GREEN_MID, C.TEAL]
        plot = chart.plots[0]
        plot.gap_width = 60
        for i, s in enumerate(plot.series):
            if kind == "line":
                s.format.line.color.rgb = rgb(colors[i % len(colors)])
                s.format.line.width = Pt(2.25)
                s.smooth = False
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
                s.format.line.fill.background()
        if number_format:
            plot.value_axis.tick_labels.number_format = number_format
            plot.value_axis.tick_labels.number_format_is_linked = False
        plot.value_axis.has_major_gridlines = True
        plot.value_axis.major_gridlines.format.line.color.rgb = rgb(C.LINE)
        plot.value_axis.format.line.color.rgb = rgb(C.LINE)
        plot.category_axis.format.line.color.rgb = rgb(C.LINE)
        return chart


# --------------------------------------------------------------------------- #
# 整套幻灯片
# --------------------------------------------------------------------------- #
class Deck:
    """整套幻灯片。"""

    def __init__(self, width=SW, height=SH):
        self.prs = Presentation()
        self.prs.slide_width = Emu(int(width * EMU))
        self.prs.slide_height = Emu(int(height * EMU))
        self.blank = self.prs.slide_layouts[6]
        self.count = 0
        self._patch_theme()

    # --- 内部 ---
    def _patch_theme(self):
        """把主题字体改成微软雅黑，避免新增占位符回落到 Calibri。"""
        try:
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            part = self.prs.slide_masters[0].part.part_related_by(RT.THEME)
            xml = part.blob.decode("utf-8")

            def repl(match):
                block = match.group(0)
                return re.sub(r'(<a:(?:latin|ea|cs) typeface=")[^"]*(")',
                              r"\g<1>微软雅黑\2", block)

            xml = re.sub(r"<a:fontScheme.*?</a:fontScheme>", repl, xml, flags=re.S)
            part._blob = xml.encode("utf-8")
        except Exception:
            pass

    def _new(self):
        self.count += 1
        return self.prs.slides.add_slide(self.blank)

    # --- 页面 ---
    def cover(self, title=None, subtitle=None, en=None, location="中国 · 北京", date=None,
              photo=None, number=True, title_color=None, overlay=0.42):
        """封面。

        * title=None：模板原味（徽标 + 公司字标 + 英文全称 + 地点）。
        * title="2026 年度市场工作会"：上方小字标 + 居中大标题 + 副标题，信息量更大。
        * photo：默认用预合成的 16:9 环保主题底图 assets/bg_cover_full.jpg。
        """
        s = self._new()
        pic(s, (0, 0, SW, SH), photo or asset("bg_cover_full.jpg"), dpi=200, name="封面底图")
        if overlay:
            rect(s, (0, 0, SW, 4.2),
                 grad=[(0, C.WHITE, overlay), (0.62, C.WHITE, overlay * 0.75),
                       (1, C.WHITE, 0.0)],
                 angle=90, name="顶部提亮")
        full_en = en or "EXAMPLE ENVIRONMENTAL GROUP CO., LTD"
        if title:
            pic(s, (6.35, 0.46, 0.62, 0.632), asset("logo_emblem.png"), name="logo")
            pic(s, (4.40, 1.16, 4.53, 0.533), asset("logo_wordmark.png"), name="公司字标")
            tsize = fit_size(title, (11.3, 0.95), 46, bold=True, line_spacing=1.15)
            text(s, (1.0, 1.92, 11.333, 0.95), title, size=tsize,
                 color=title_color or C.INK_DEEP, bold=True, align="c", anchor="m")
            rect(s, (5.67, 2.97, 2.0, 0.06),
                 grad=[(0, C.MINT, None), (1, C.GREEN, None)], angle=0, name="装饰线")
            if subtitle:
                text(s, (1.5, 3.12, 10.33, 0.42), subtitle, size=17, color=C.INK, align="c")
            if en:
                text(s, (2.5, 3.58, 8.333, 0.30), en, size=11, color=C.GRAY, align="c",
                     spacing=1.2)
        else:
            pic(s, (6.20, 0.962, 0.93, 0.948), asset("logo_emblem.png"), name="logo")
            pic(s, (2.58, 2.49, 8.55, 1.006), asset("logo_wordmark.png"), name="公司字标")
            text(s, (2.71, 3.34, 8.21, 0.34), full_en, size=14, color="606060", align="c",
                 spacing=1.5)
        bottom = "    ".join([v for v in [location, date] if v])
        if bottom:
            text(s, (2.0, 6.48, 9.333, 0.52), bottom, size=24, color=C.WHITE, bold=True,
                 align="c", anchor="m", text_shadow=True, spacing=2)
        if number:
            page_number(s, 1, color="C9D8E8")
        return s

    def toc(self, items, title="目录", en="CONTENTS", subtitle=None, number=True):
        """目录页：左侧标题 + 右侧编号卡片，背后两条大箭头色带。"""
        s = self._new()
        pic(s, (0, 0, SW, SH), asset("bg_light.jpg"), dpi=200, name="背景")
        _chev(s, (4.54, 0, 6.76, 7.50), None,
              [(0, "50C563", 0.16), (1, C.GREEN_MID, None)], 90, 0.5, None, "色带")
        _chev(s, (6.36, 0.70, 5.50, 6.10), None,
              [(0, C.GREEN_SOFT, None), (0.97, C.TEAL, 0.45)], 140, 0.5, None, "色带")

        # 左侧标题块
        ty = 2.38
        if title == "目录":
            pic(s, (1.20, ty, 1.59, 1.007), asset("wordmark_mulu.png"), name="目录字标")
        else:
            text(s, (1.20, ty, 3.6, 1.0), title, size=46, color=C.INK, bold=True, anchor="m")
        text(s, (1.20, ty + 1.10, 2.90, 0.50), en, size=28, color=C.GRAY_LIGHT,
             spacing=6, align="l", anchor="m")
        rect(s, (1.35, ty + 1.66, 2.44, 0.12),
             grad=[(0, C.GREEN_SOFT, None), (1, "09E4D8", None)], angle=140,
             shape=MSO_SHAPE.PARALLELOGRAM, name="标题装饰块")
        if subtitle:
            text(s, (1.20, ty + 1.92, 3.2, 0.6), subtitle, size=12, color=C.GRAY,
                 line_spacing=1.3)

        # 右侧卡片
        n = len(items)
        card_x, card_w, card_h = 4.973, 6.427, 0.93
        gap = 0.307 if n <= 4 else 0.18
        total = n * card_h + (n - 1) * gap
        y0 = (SH - total) / 2 + 0.05
        for i, item in enumerate(items):
            if isinstance(item, (tuple, list)):
                name, sub = item[0], (item[1] if len(item) > 1 else None)
            else:
                name, sub = item, None
            cy = y0 + i * (card_h + gap)
            rect(s, (card_x, cy, card_w, card_h), radius=0.155,
                 grad=[(0, "9FC4EE", None), (0.5, "DCEAF8", None), (1, C.WHITE, None)],
                 angle=140,
                 shadow_kw=dict(blur_pt=16, color="15329B", alpha=0.40, glow=True),
                 name="目录卡片")
            text(s, (card_x + 0.418, cy + 0.06, 0.64, 0.52), f"{i + 1:02d}", size=28,
                 color=C.BLUE, bold=True, align="c", anchor="m")
            for k in range(4):
                _chev(s, (card_x + 0.528 + k * 0.2085, cy + 0.60, 0.185, 0.238),
                      "C9DCF0", None, 0, 0.504, None, "小箭头")
            tsize = fit_size(name, (card_w - 2.0, card_h), 32, bold=True, line_spacing=1.1)
            paras = [[(name, {"size": tsize, "bold": True, "color": C.INK})]]
            if sub:
                paras.append([(sub, {"size": 12, "color": C.GRAY})])
            text(s, (card_x + 1.404, cy, card_w - 1.6, card_h), paras,
                 anchor="m", line_spacing=1.1)
        logo(s)
        if number:
            page_number(s, self.count)
        return s

    def section(self, no, title, en=None, photo=None, caption=None, number=True,
                layout="left", focus=(0.5, 0.5)):
        """章节过渡页：左侧大编号 + 标题，右侧满幅照片，最左绿色尖角。"""
        s = self._new()
        rect(s, (0, 0, SW, SH), color="F7FAFD", name="背景")
        if photo:
            if layout == "left":
                pic(s, (0, 0, SW, SH), photo, dpi=200, focus=focus, name="章节配图")
                rect(s, (0, 0, 7.9, SH),
                     grad=[(0, C.WHITE, 0.98), (0.55, C.WHITE, 0.80), (1, C.WHITE, 0.0)],
                     angle=0, name="左侧留白")
            else:
                pic(s, (6.35, 0, 6.983, SH), photo, dpi=200, focus=focus, name="章节配图")
                rect(s, (5.35, 0, 2.4, SH),
                     grad=[(0, C.WHITE, 1.0), (1, C.WHITE, 0.0)], angle=0, name="左侧留白")
        rect(s, (0, 2.445, 0.586, 1.827), color=C.GREEN,
             shape=MSO_SHAPE.PENTAGON, name="绿色尖角")
        text(s, (1.06, 1.66, 2.3, 1.62), str(no), size=84, color=C.GRAY_LIGHT, bold=True,
             align="l", anchor="m", line_spacing=1.0)
        tsize = fit_size(title, (6.0, 1.0), 44, bold=True, line_spacing=1.15)
        text(s, (1.12, 3.32, 6.1, 1.0), title, size=tsize, color=C.INK, bold=True,
             anchor="m")
        cursor = 4.36
        if en:
            text(s, (1.14, cursor, 5.6, 0.30), en, size=12, color=C.GRAY, spacing=2.5)
            cursor += 0.34
        rect(s, (1.14, cursor, 1.30, 0.06),
             grad=[(0, C.MINT, None), (1, C.GREEN, None)], angle=0)
        if caption:
            text(s, (1.14, cursor + 0.16, 5.4, 0.7), caption, size=13, color=C.INK,
                 line_spacing=1.35)
        logo(s)
        if number:
            page_number(s, self.count, color=C.GRAY)
        return s

    def page(self, title, subtitle=None, number=True, wave=True, logo_on=True,
             box=CONTENT_BOX):
        """标准内容页。返回 Page 对象，用它摆内容块。"""
        s = self._new()
        pic(s, (0, 0, SW, SH), asset("bg_light.jpg"), dpi=200, name="背景")
        if wave:
            pts = [(0, 6.92), (5.35, SH), (0, SH)]
            ff = s.shapes.build_freeform(Emu(int(pts[0][0] * EMU)),
                                         Emu(int(pts[0][1] * EMU)))
            ff.add_line_segments([(Emu(int(px * EMU)), Emu(int(py * EMU)))
                                  for px, py in pts[1:]], close=True)
            wave_shape = ff.convert_to_shape()
            fill_grad(wave_shape, [(0, C.MINT, 0.75), (1, C.GREEN, 0.95)], 135)
            no_line(wave_shape)
            wave_shape.name = "底部绿带"
        tiny_chevrons(s, 0.0, 0.487, 4, h=0.238, w=0.185,
                      colors=(C.MINT_PALE, C.GREEN))
        tsize = fit_size(title, (10.3, 0.60), 28, bold=True, line_spacing=1.1)
        text(s, (0.800, 0.30, 10.4, 0.62), title, size=tsize, color=C.INK, bold=True,
             anchor="m")
        if subtitle:
            text(s, (0.800, 0.92, 10.6, 0.30), subtitle, size=12, color=C.GRAY)
        if logo_on:
            logo(s)
        if number:
            page_number(s, self.count)
        return Page(s, title, box)

    def closing(self, thanks="分享结束 感谢聆听", company="示例环保集团",
                address="XX 市 XX 区 XX 路 XX 号",
                en="THANK YOU FOR LISTENING", photo=None, qr=True, number=True,
                thanks_image=True):
        """结尾页：致辞图 + 英文 + 底部绿色信息条 + 二维码。"""
        s = self._new()
        pic(s, (0, 0, SW, SH), photo or asset("bg_closing.jpg"), dpi=200, name="结尾底图")
        if thanks_image and thanks == "分享结束 感谢聆听":
            pic(s, (3.75, 2.43, 5.83, 0.945), asset("wordmark_thanks.png"), name="致辞字标")
        else:
            text(s, (2.0, 2.50, 9.333, 0.85), thanks, size=40, color=C.GREEN_DEEP,
                 bold=True, align="c", anchor="m", text_shadow=True)
        text(s, (2.6, 3.52, 8.133, 0.34), en, size=14, color="606060", align="c", spacing=1.6)
        rect(s, (0, 6.27, SW, 1.23),
             grad=[(0, "4CB268", None), (1, "2F8E4E", None)], angle=0, name="信息条")
        if qr:
            rect(s, (2.22, 5.77, 1.41, 1.41), color=C.WHITE, radius=0.09, name="二维码底板")
            pic(s, (2.31, 5.86, 1.23, 1.23), asset("qr_wechat.png"), dpi=200, name="二维码")
        text(s, (4.23, 6.40, 7.6, 0.36), company, size=20, color=C.WHITE, bold=True,
             anchor="m")
        text(s, (4.51, 6.80, 8.3, 0.55), address, size=13, color=C.WHITE, anchor="t",
             line_spacing=1.25)
        if number:
            page_number(s, self.count, color=C.WHITE)
        return s

    # --- 输出 ---
    def save(self, path):
        out = Path(path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return out
