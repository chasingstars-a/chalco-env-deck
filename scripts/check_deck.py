# -*- coding: utf-8 -*-
"""成品自检：尺寸 / 字体 / 配色 / logo / 页码 / 文字溢出 / 越界。

    python check_deck.py ..\\demo\\示例环保-公司简介-示例.pptx

退出码 0 = 没有 error（warning 仅提示）。
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from chalco_deck import FONT, PALETTE, SH, SW, asset, wrap_lines  # noqa: E402

EMU = 914400
FORBIDDEN = {"FF0000", "C00000", "E60000", "CC0000", "FF3333"}   # 旧版红色强调

errors: list[str] = []
warnings: list[str] = []
infos: list[str] = []


def inch(v):
    return round(v / EMU, 3)


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:                     # GROUP
            yield from walk(sh.shapes)


def runs_of(tf):
    for p in tf.paragraphs:
        for r in p.runs:
            if r.text.strip():
                yield p, r


def iter_text_shapes(slide):
    for sh in walk(slide.shapes):
        if sh.has_text_frame and sh.text.strip():
            yield sh


def check_fonts(slide_no, slide):
    for sh in iter_text_shapes(slide):
        for p, r in runs_of(sh.text_frame):
            rPr = r._r.find(qn("a:rPr"))
            if rPr is None:
                continue
            for tag in ("a:latin", "a:ea"):
                el = rPr.find(qn(tag))
                face = None if el is None else el.get("typeface")
                if face and face != FONT:
                    warnings.append(f"第 {slide_no} 页 「{sh.text[:12]}」字体为 {face}，应为 {FONT}")


def check_colors(slide_no, slide):
    for sh in walk(slide.shapes):
        el = sh._element
        for clr in el.iter(qn("a:srgbClr")):
            value = (clr.get("val") or "").upper()
            if value in FORBIDDEN:
                errors.append(f"第 {slide_no} 页 出现旧版红色强调 #{value}")
            elif value and value not in PALETTE:
                warnings.append(f"第 {slide_no} 页 出现调色板外颜色 #{value}")


def check_overflow(slide_no, slide):
    for sh in iter_text_shapes(slide):
        tf = sh.text_frame
        box_w = inch(sh.width) - inch(tf.margin_left or 0) - inch(tf.margin_right or 0)
        box_h = inch(sh.height) - inch(tf.margin_top or 0) - inch(tf.margin_bottom or 0)
        if box_w <= 0.2 or box_h <= 0.1:
            continue
        need = 0.0
        for p in tf.paragraphs:
            runs = [r for r in p.runs if r.text]
            if not runs:
                continue
            size = max((r.font.size.pt if r.font.size else 18) for r in runs)
            bold = any(bool(r.font.bold) for r in runs)
            text = "".join(r.text for r in runs)
            ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
            need += len(wrap_lines(text, box_w, size, bold)) * size * 1.2 * ls / 72
            if p.space_after is not None:
                need += p.space_after.pt / 72
        if need > box_h * 1.08 + 0.06:
            label = sh.text.replace("\n", " ")[:18]
            warnings.append(
                f"第 {slide_no} 页 「{label}」文字可能溢出："
                f"需要 {need:.2f}in / 框高 {box_h:.2f}in")
        if sh.left < 0 or sh.top < 0 or inch(sh.left) + inch(sh.width) > SW + 0.02 \
                or inch(sh.top) + inch(sh.height) > SH + 0.02:
            warnings.append(f"第 {slide_no} 页 「{sh.text[:12]}」文字框超出页面边界")


def check_logo(slide_no, slide, logo_hashes, expect):
    found = False
    for sh in walk(slide.shapes):
        if sh.shape_type == 13:
            if sh.name == "logo":            # chalco_deck.logo() 的命名约定
                found = True
                break
            try:
                blob = sh.image.blob
            except Exception:
                continue
            if hashlib.sha1(blob).hexdigest() in logo_hashes:
                found = True
                break
    if expect and not found:
        errors.append(f"第 {slide_no} 页 缺少 EXAMPLE 徽标")
    if not expect and not found:
        infos.append(f"第 {slide_no} 页 未放徽标（封面/章节页请确认是否为有意）")


def check_page_number(slide_no, slide, total):
    hits = []
    for sh in iter_text_shapes(slide):
        text = sh.text.strip()
        if text.isdigit() and int(text) == slide_no:
            if inch(sh.left) > SW * 0.85 and inch(sh.top) > SH * 0.85:
                hits.append(text)
    if slide_no < total and not hits:
        warnings.append(f"第 {slide_no} 页 右下角缺少页码")


def main(paths):
    for path in paths:
        path = Path(path)
        prs = Presentation(str(path))
        total = len(prs.slides)
        if abs(inch(prs.slide_width) - SW) > 0.02 or abs(inch(prs.slide_height) - SH) > 0.02:
            errors.append(f"页面尺寸 {inch(prs.slide_width)}x{inch(prs.slide_height)}，"
                          f"应为 {SW}x{SH}（16:9）")
        logo_hashes = set()
        for name in ("logo_emblem.png", "logo_emblem_white.png"):
            logo_hashes.add(hashlib.sha1(Path(asset(name)).read_bytes()).hexdigest())

        print(f"\n=== {path.name}：{total} 页 ===")
        for i, slide in enumerate(prs.slides, start=1):
            check_fonts(i, slide)
            check_colors(i, slide)
            check_overflow(i, slide)
            expect_logo = 1 < i < total
            check_logo(i, slide, logo_hashes, expect_logo)
            check_page_number(i, slide, total)
            title = ""
            for sh in iter_text_shapes(slide):
                if sh.height and inch(sh.top) < 0.9 and len(sh.text) < 30:
                    title = sh.text
                    break
            print(f"  {i:02d} {title[:30]:32s} shapes={len(slide.shapes)}")

    if warnings:
        print("\n-- 警告 --")
        for w in warnings:
            print("  ⚠", w)
    if errors:
        print("\n-- 错误 --")
        for e in errors:
            print("  ✗", e)
    if infos:
        print("\n-- 提示 --")
        for i in infos[:10]:
            print("  ·", i)
    print(f"\n结果：{len(errors)} 个错误，{len(warnings)} 个警告")
    return 1 if errors else 0


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(args))
